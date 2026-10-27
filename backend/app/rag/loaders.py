"""Turn PDFs / DOCX / HTML / Markdown / URLs into Documents made of heading-labelled sections."""
import hashlib
import json
import re
from pathlib import Path

import httpx
from bs4 import BeautifulSoup

from .models import Document, Section

SUPPORTED = {".pdf", ".docx", ".md", ".txt", ".html", ".htm"}


def _doc_id(key: str) -> str:
    return hashlib.sha1(key.encode()).hexdigest()[:12]


def _title_from_name(name: str) -> str:
    return re.sub(r"[_\-]+", " ", Path(name).stem).strip().title()


def load_manifest(docs_dir: Path) -> dict:
    p = docs_dir / "manifest.json"
    if p.exists():
        return json.loads(p.read_text())
    return {}


def _split_markdown(text: str) -> list[Section]:
    sections: list[Section] = []
    path: list[tuple[int, str]] = []
    buf: list[str] = []

    def flush():
        body = "\n".join(buf).strip()
        if body:
            sections.append(Section(" > ".join(h for _, h in path), body))
        buf.clear()

    for line in text.splitlines():
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            flush()
            level = len(m.group(1))
            while path and path[-1][0] >= level:
                path.pop()
            path.append((level, m.group(2).strip()))
        else:
            buf.append(line)
    flush()
    return sections


def _load_pdf(p: Path) -> list[Section]:
    from pypdf import PdfReader
    out = []
    for i, page in enumerate(PdfReader(str(p)).pages, start=1):
        t = (page.extract_text() or "").strip()
        if t:

            first = t.splitlines()[0][:100]
            out.append(Section(first if len(first) < 90 else "", t, page=i))
    return out


def _load_docx(p: Path) -> list[Section]:
    from docx import Document as Docx
    d = Docx(str(p))
    sections: list[Section] = []
    heading, buf = "", []

    def flush():
        body = "\n".join(buf).strip()
        if body:
            sections.append(Section(heading, body))
        buf.clear()

    for para in d.paragraphs:
        style = (para.style.name or "").lower() if para.style is not None else ""
        if style.startswith("heading") and para.text.strip():
            flush()
            heading = para.text.strip()
        elif para.text.strip():
            buf.append(para.text.strip())
    for table in d.tables:
        for row in table.rows:
            buf.append(" | ".join(c.text.strip() for c in row.cells))
    flush()
    return sections


def _html_to_sections(html: str) -> tuple[str, list[Section]]:
    soup = BeautifulSoup(html, "html.parser")
    for t in soup(["script", "style", "nav", "footer", "header", "noscript"]):
        t.decompose()
    title = (soup.title.string.strip() if soup.title and soup.title.string else "")
    sections: list[Section] = []
    heading, buf = title, []

    def flush():
        body = "\n".join(buf).strip()
        if body:
            sections.append(Section(heading, body))
        buf.clear()

    root = soup.find("main") or soup.body or soup
    for el in root.find_all(["h1", "h2", "h3", "h4", "p", "li", "tr"]):
        txt = el.get_text(" ", strip=True)
        if not txt:
            continue
        if el.name in {"h1", "h2", "h3", "h4"}:
            flush()
            heading = txt
        else:
            buf.append(txt)
    flush()
    return title, sections


def load_file(p: Path, meta: dict | None = None) -> Document | None:
    ext = p.suffix.lower()
    if ext not in SUPPORTED:
        return None
    meta = meta or {}
    title, sections = meta.get("title") or _title_from_name(p.name), []
    if ext == ".pdf":
        sections = _load_pdf(p)
    elif ext == ".docx":
        sections = _load_docx(p)
    elif ext in {".md", ".txt"}:
        text = p.read_text(errors="ignore")
        m = re.match(r"^#\s+(.*)$", text.strip().splitlines()[0]) if text.strip() else None
        if m and not meta.get("title"):
            title = m.group(1).strip()
        sections = _split_markdown(text)
    else:
        t, sections = _html_to_sections(p.read_text(errors="ignore"))
        if t and not meta.get("title"):
            title = t
    if not sections:
        return None
    return Document(
        doc_id=_doc_id(p.name), title=title, url=meta.get("url", ""), doc_type=meta.get("doc_type", "guidance"),
        filename=p.name, sections=sections,
    )


def load_url(url: str, meta: dict | None = None) -> Document | None:
    meta = meta or {}
    r = httpx.get(url, follow_redirects=True, timeout=30, headers={"User-Agent": "EsolvIngest/0.1"})
    r.raise_for_status()
    ctype = r.headers.get("content-type", "")
    if "pdf" in ctype or url.lower().endswith(".pdf"):
        import io
        from pypdf import PdfReader
        sections = []
        for i, page in enumerate(PdfReader(io.BytesIO(r.content)).pages, start=1):
            t = (page.extract_text() or "").strip()
            if t:
                sections.append(Section("", t, page=i))
        title = meta.get("title") or _title_from_name(url.rsplit("/", 1)[-1])
    else:
        title, sections = _html_to_sections(r.text)
        title = meta.get("title") or title or url
    if not sections:
        return None
    return Document(doc_id=_doc_id(url), title=title, url=url, doc_type=meta.get("doc_type", "guidance"), sections=sections)


def load_all(docs_dir: Path, sources_file: Path | None = None) -> list[Document]:
    manifest = load_manifest(docs_dir)
    docs: list[Document] = []
    for p in sorted(docs_dir.rglob("*")):
        if p.is_file() and p.name != "manifest.json" and not p.name.startswith("."):
            meta = manifest.get(p.name, {})
            if meta.get("index") is False:
                continue
            d = load_file(p, meta)
            if d:
                docs.append(d)
    if sources_file and sources_file.exists():
        for line in sources_file.read_text().splitlines():
            line = line.strip()
            if not line or line.startswith("#"):
                continue
            url, _, doc_type = line.partition("|")
            try:
                d = load_url(url.strip(), {"doc_type": doc_type.strip() or "guidance"})
                if d:
                    docs.append(d)
            except Exception as e:
                print(f"! failed to fetch {url}: {e}")
    return docs
