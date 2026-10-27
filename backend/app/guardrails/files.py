"""Teacher-attached documents: validate, extract plain text, cap size. Nothing is stored; the text rides along with one chat request."""
import io
import re
import zipfile
from pathlib import Path

from bs4 import BeautifulSoup

from ..config import get_settings

ALLOWED = {".pdf", ".docx", ".txt", ".md", ".csv", ".html", ".htm"}
MAX_PDF_PAGES = 60
MAX_DOCX_UNZIPPED = 60 * 1024 * 1024


class FileRejected(ValueError):
    pass


def _pdf(raw: bytes) -> tuple[str, int]:
    from pypdf import PdfReader
    try:
        reader = PdfReader(io.BytesIO(raw))
        if reader.is_encrypted:
            raise FileRejected("That PDF is password-protected.")
        pages = reader.pages[:MAX_PDF_PAGES]
        return "\n\n".join((p.extract_text() or "").strip() for p in pages), len(reader.pages)
    except FileRejected:
        raise
    except Exception as e:
        raise FileRejected("That PDF couldn't be read.") from e


def _docx(raw: bytes) -> tuple[str, int]:
    from docx import Document
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            if sum(i.file_size for i in z.infolist()) > MAX_DOCX_UNZIPPED:
                raise FileRejected("That document is too large to read.")
        d = Document(io.BytesIO(raw))
    except FileRejected:
        raise
    except Exception as e:
        raise FileRejected("That Word document couldn't be read.") from e
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts), 0


def _html(raw: bytes) -> tuple[str, int]:
    soup = BeautifulSoup(raw.decode("utf-8", "ignore"), "html.parser")
    for t in soup(["script", "style", "noscript"]):
        t.decompose()
    return soup.get_text("\n", strip=True), 0


def extract(filename: str, raw: bytes) -> dict:
    """Return {name, text, chars, truncated, pages}. Raises FileRejected with a teacher-friendly message."""
    s = get_settings()
    name = Path(filename or "file").name[:120]
    ext = Path(name).suffix.lower()
    if ext not in ALLOWED:
        raise FileRejected("Supported files: PDF, Word (.docx), text, Markdown, CSV and HTML. For pictures, attach them as images.")
    if len(raw) > s.max_file_bytes:
        raise FileRejected(f"Files must be under {s.max_file_bytes // (1024 * 1024)} MB.")
    if not raw:
        raise FileRejected("That file is empty.")

    if ext == ".pdf" and not raw.lstrip().startswith(b"%PDF"):
        raise FileRejected("That doesn't look like a real PDF.")
    if ext == ".docx" and not raw.startswith(b"PK"):
        raise FileRejected("That doesn't look like a real Word document.")
    pages = 0
    if ext == ".pdf":
        text, pages = _pdf(raw)
    elif ext == ".docx":
        text, pages = _docx(raw)
    elif ext in {".html", ".htm"}:
        text, pages = _html(raw)
    else:
        text = raw.decode("utf-8", "ignore")
    text = re.sub(r"[ \t]+\n", "\n", re.sub(r"\n{3,}", "\n\n", text)).strip()
    if len(text) < 5:
        raise FileRejected("I couldn't find readable text in that file. If it's a scan or photo, attach it as an image instead.")
    chars = len(text)
    truncated = chars > s.max_file_chars
    return {"name": name, "text": text[: s.max_file_chars], "chars": chars, "truncated": truncated, "pages": pages}
